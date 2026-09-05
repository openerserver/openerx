from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / "tests" / "v2" / "fixtures" / "m4"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def create_pdf() -> None:
    output = FIXTURES / "file.cited-pdf.v1.pdf"
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            "FixtureTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=25,
            leading=30,
            textColor=HexColor("#14372B"),
            spaceAfter=18,
        )
    )
    styles.add(
        ParagraphStyle(
            "FixtureBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=11,
            leading=17,
            textColor=HexColor("#27342F"),
        )
    )
    document = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        rightMargin=22 * mm,
        leftMargin=22 * mm,
        topMargin=22 * mm,
        bottomMargin=22 * mm,
        title="OpenERX M4 cited PDF fixture",
        author="OpenERX M4 tests",
    )
    story = [
        Paragraph("M4 File Grounding Brief", styles["FixtureTitle"]),
        Paragraph("Page 1 establishes the fixture and citation rules.", styles["FixtureBody"]),
        Spacer(1, 8 * mm),
        Table(
            [["Fixture", "file.cited-pdf.v1"], ["Expected citation", "Page 2"]],
            colWidths=[42 * mm, 102 * mm],
            style=TableStyle(
                [
                    ("BACKGROUND", (0, 0), (0, -1), HexColor("#DCEBE3")),
                    ("TEXTCOLOR", (0, 0), (-1, -1), HexColor("#27342F")),
                    ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.5, HexColor("#9FB7AA")),
                    ("PADDING", (0, 0), (-1, -1), 8),
                ]
            ),
        ),
        PageBreak(),
        Paragraph("The decision is on this page", styles["FixtureTitle"]),
        Paragraph(
            "GT-FILE-01 PAGE TWO FACT: The launch window is 09:30. "
            "This sentence must be grounded with a page 2 citation.",
            styles["FixtureBody"],
        ),
        Spacer(1, 10 * mm),
        Paragraph(
            "The source remains the original PDF while OpenERX works from a controlled copy.",
            styles["FixtureBody"],
        ),
    ]
    document.build(story)


def create_docx() -> None:
    output = FIXTURES / "file.docx-pair.v1.docx"
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(39, 52, 47)
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.color.rgb = RGBColor(20, 55, 43)

    title = document.add_paragraph(style="Title")
    title.add_run("M4 Approval Memo")
    subtitle = document.add_paragraph("Controlled document fixture · file.docx-pair.v1")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle.runs[0].font.color.rgb = RGBColor(83, 108, 96)

    document.add_heading("Decision", level=1)
    document.add_paragraph(
        "Approve the file and artifact alpha when every hard gate is backed by reproducible evidence."
    )
    document.add_heading("Requested changes", level=1)
    table = document.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    headers = ["Item", "Current", "Approved"]
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
        set_cell_shading(table.rows[0].cells[index], "DCEBE3")
    rows = [
        ("Scope", "Files only", "Files and folders"),
        ("Versions", "Replace", "Append immutable version"),
        ("Citation", "Text only", "Page, sheet/range, slide, text range"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value

    document.add_page_break()
    document.add_heading("Fixed verification text", level=1)
    document.add_paragraph(
        "GT-FILE-06 FIXED TEXT: Originals are preserved; approved edits create a new artifact version."
    )
    document.add_heading("Acceptance", level=2)
    for text in [
        "The document opens and renders without clipped text.",
        "Heading styles remain present.",
        "The approved version does not overwrite the source.",
    ]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(7)
        bullet = paragraph.add_run("•  ")
        bullet.font.color.rgb = RGBColor(104, 165, 126)
        paragraph.add_run(text)

    footer = section.footer
    footer_table = footer.add_table(rows=1, cols=2, width=Inches(6.8))
    footer_table.rows[0].cells[0].text = "OpenERX M4"
    footer_table.rows[0].cells[1].text = "Controlled fixture"
    footer_table.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.core_properties.title = "OpenERX M4 Approval Memo"
    document.core_properties.author = "OpenERX M4 tests"
    document.save(output)


def create_ocr_image() -> None:
    output = FIXTURES / "file.ocr-image.v1.png"
    image = Image.new("RGB", (1200, 520), "#F3F6F2")
    draw = ImageDraw.Draw(image)
    try:
        title_font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 54)
        body_font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 42)
    except OSError:
        title_font = ImageFont.load_default()
        body_font = ImageFont.load_default()
    draw.rounded_rectangle((55, 55, 1145, 465), radius=30, fill="#FFFFFF", outline="#9FB7AA", width=4)
    draw.text((105, 105), "M4 OCR CHECK", font=title_font, fill="#14372B")
    draw.text((105, 225), "Archive box: 184?", font=body_font, fill="#27342F")
    draw.text((105, 315), "The final character is intentionally uncertain.", font=body_font, fill="#5B6D64")
    image.save(output)


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    create_pdf()
    create_docx()
    create_ocr_image()


if __name__ == "__main__":
    main()
